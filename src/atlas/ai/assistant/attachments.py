"""Chat attachments: local-only parsing into text/tables. Untrusted by construction.

* Allowed types by EXTENSION *and* magic bytes: csv/tsv/txt, pdf, xlsx, png/jpg/gif/webp, xml (ISO 20022 camt/pain),
  plus broker statements in those formats. Anything else is refused (no archives, no executables, no macro workbooks).
* Size limits per type; a total cap per conversation.
* Stored in a quarantine dir (0700) with 0600 files under a random id; original names are never used as paths.
* Nothing is executed: PDF via `pdftotext` (argv list, no shell, timeout, no network), XLSX via openpyxl read-only
  data_only (formulas/macros never evaluated), XML via defusedxml (XXE / billion-laughs / external entities refused),
  images via Pillow metadata only (+ local `tesseract` OCR when installed).
* Extracted text is PII-redacted, injection-flagged and always fenced as untrusted when given to a model.
* Retention: files are deleted after `retention_days` or on request.
"""
from __future__ import annotations

import csv
import io
import json
import os
import secrets
import shutil
import stat
import subprocess
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path

from atlas.ai.assistant.guard import injection_flags, redact

MB = 1024 * 1024
LIMITS = {"csv": 5 * MB, "tsv": 5 * MB, "txt": 2 * MB, "pdf": 15 * MB, "xlsx": 10 * MB, "xml": 5 * MB,
          "png": 8 * MB, "jpg": 8 * MB, "jpeg": 8 * MB, "gif": 4 * MB, "webp": 8 * MB}
MAX_TOTAL_PER_CONV = 40 * MB
MAGIC = {"pdf": [b"%PDF-"], "xlsx": [b"PK\x03\x04"], "png": [b"\x89PNG\r\n\x1a\n"], "jpg": [b"\xff\xd8\xff"], "jpeg": [b"\xff\xd8\xff"],
         "gif": [b"GIF87a", b"GIF89a"], "webp": [b"RIFF"]}
TEXTLIKE = {"csv", "tsv", "txt", "xml"}
MAX_ROWS, MAX_TEXT = 2000, 60_000


class AttachmentError(ValueError):
    def __init__(self, code: str, msg: str):
        super().__init__(msg)
        self.code = code


@dataclass
class Attachment:
    id: str
    name: str
    ext: str
    size: int
    sha256: str
    created: float
    kind: str = "unknown"            # table | text | iso20022 | broker_statement | image
    summary: str = ""
    text: str = ""
    table: dict | None = None        # {"columns": [...], "rows": [[...]], "truncated": bool}
    meta: dict = field(default_factory=dict)
    redactions: dict = field(default_factory=dict)
    injection: list = field(default_factory=list)
    owner: str = ""

    def public(self) -> dict:
        d = asdict(self)
        d["text"] = self.text[:2000]
        if self.table:
            d["table"] = {**self.table, "rows": self.table["rows"][:50]}
        return d


def _sniff_ext(name: str) -> str:
    base = os.path.basename(name or "")
    if "." not in base:
        raise AttachmentError("unsupported", "file has no extension")
    ext = base.rsplit(".", 1)[1].lower()
    if ext not in LIMITS:
        raise AttachmentError("unsupported", f".{ext} files are not supported (allowed: {', '.join(sorted(LIMITS))})")
    return ext


def _check_magic(ext: str, data: bytes) -> None:
    if ext in MAGIC and not any(data.startswith(m) for m in MAGIC[ext]):
        raise AttachmentError("unsupported", f"content does not look like a .{ext} file")
    if ext == "webp" and data[8:12] != b"WEBP":
        raise AttachmentError("unsupported", "content does not look like a .webp file")
    if ext in TEXTLIKE:
        if b"\x00" in data[:4096]:
            raise AttachmentError("unsupported", "binary content in a text file")
        if data.startswith((b"MZ", b"\x7fELF", b"PK\x03\x04", b"#!")):
            raise AttachmentError("unsupported", "executable or archive content refused")


class AttachmentStore:
    def __init__(self, root: str | Path, retention_days: int = 30, ocr: bool | None = None):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        os.chmod(self.root, 0o700)
        self.retention_s = retention_days * 86400
        self.ocr = shutil.which("tesseract") is not None if ocr is None else ocr

    # ---------------------------------------------------------------- lifecycle
    def _path(self, aid: str) -> Path:
        if not aid.startswith("att_") or not aid[4:].isalnum():
            raise AttachmentError("not_found", "bad attachment id")
        return self.root / aid

    def save(self, owner: str, name: str, data: bytes, conv_bytes: int = 0) -> Attachment:
        import hashlib
        ext = _sniff_ext(name)
        if not data:
            raise AttachmentError("empty", "file is empty")
        if len(data) > LIMITS[ext]:
            raise AttachmentError("too_large", f".{ext} files are limited to {LIMITS[ext] // MB} MB (this one is {len(data) / MB:.1f} MB)")
        if conv_bytes + len(data) > MAX_TOTAL_PER_CONV:
            raise AttachmentError("too_large", f"attachments in one conversation are limited to {MAX_TOTAL_PER_CONV // MB} MB")
        _check_magic(ext, data)
        aid = "att_" + secrets.token_hex(8)
        d = self._path(aid)
        d.mkdir(mode=0o700)
        raw = d / f"blob.{ext}"                  # never the user's filename
        fd = os.open(raw, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        safe_name = "".join(c for c in os.path.basename(name) if c.isprintable())[:120]
        a = Attachment(aid, safe_name, ext, len(data), hashlib.sha256(data).hexdigest(), time.time(), owner=owner)
        try:
            self._parse(a, raw, data)
        except AttachmentError:
            shutil.rmtree(d, ignore_errors=True)
            raise
        except Exception as e:  # noqa: BLE001 - a parser failure is reported, never crashes the chat
            shutil.rmtree(d, ignore_errors=True)
            raise AttachmentError("parse_failed", f"could not read this .{ext} file ({type(e).__name__})")
        a.text, red = redact(a.text)
        a.redactions = red
        if a.table:
            rows = []
            for r in a.table["rows"]:
                rr = [redact(str(c))[0] for c in r]
                rows.append(rr)
            a.table["rows"] = rows
        a.injection = injection_flags(a.text + " " + json.dumps(a.table or {}))
        meta = d / "meta.json"
        fd = os.open(meta, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "w") as f:
            json.dump(asdict(a), f)
        return a

    def get(self, aid: str, owner: str | None = None) -> Attachment:
        p = self._path(aid) / "meta.json"
        if not p.exists():
            raise AttachmentError("not_found", "attachment not found (deleted or expired)")
        a = Attachment(**json.loads(p.read_text()))
        if owner is not None and a.owner != owner:
            raise AttachmentError("not_found", "attachment not found")
        return a

    def delete(self, aid: str, owner: str | None = None) -> bool:
        self.get(aid, owner)
        shutil.rmtree(self._path(aid), ignore_errors=True)
        return True

    def purge_expired(self, now: float | None = None) -> int:
        now = now or time.time()
        n = 0
        for d in self.root.glob("att_*"):
            m = d / "meta.json"
            try:
                created = json.loads(m.read_text())["created"]
            except Exception:  # noqa: BLE001
                created = d.stat().st_mtime
            if now - created > self.retention_s:
                shutil.rmtree(d, ignore_errors=True); n += 1
        return n

    @staticmethod
    def perms_ok(path: Path) -> bool:
        return stat.S_IMODE(path.stat().st_mode) & 0o077 == 0

    # ---------------------------------------------------------------- parsers
    def _parse(self, a: Attachment, path: Path, data: bytes) -> None:
        if a.ext in ("csv", "tsv", "txt"):
            text = data.decode("utf-8-sig", errors="replace")
            if a.ext == "txt":
                a.kind, a.text, a.summary = "text", text[:MAX_TEXT], f"text, {len(text.splitlines())} lines"
                return
            self._table_from_rows(a, list(csv.reader(io.StringIO(text), delimiter="\t" if a.ext == "tsv" else
                                                       csv.Sniffer().sniff(text[:4096], ",;|\t").delimiter if text.strip() else ",")))
        elif a.ext == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False, keep_links=False)
            ws = wb.worksheets[0]
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS:
                    break
                rows.append(["" if v is None else str(v) for v in r])
            a.meta["sheets"] = wb.sheetnames[:20]
            wb.close()
            self._table_from_rows(a, rows)
        elif a.ext == "pdf":
            try:
                out = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", "-l", "50", str(path), "-"], capture_output=True,
                                     timeout=20, check=False, env={"PATH": os.environ.get("PATH", "/usr/bin")})
                if out.returncode != 0:
                    raise AttachmentError("parse_failed", "PDF could not be read (encrypted or damaged?)")
                text = out.stdout.decode("utf-8", errors="replace")[:MAX_TEXT]
            except subprocess.TimeoutExpired:
                raise AttachmentError("parse_failed", "PDF took too long to read (over 20 s)")
            except FileNotFoundError:
                import re
                if b"%%EOF" not in data or b"endobj" not in data:
                    raise AttachmentError("parse_failed", "PDF could not be read (encrypted or damaged?)")
                raw_txt = data.decode("latin-1", errors="replace")
                chunks = []
                for sm in re.finditer(r"stream\r?\n(.*?)\r?\nendstream", raw_txt, re.DOTALL):
                    chunks.extend(re.findall(r"\(((?:\\.|[^\\()])*)\)", sm.group(1)))
                text = "\n".join(c.replace("\\(", "(").replace("\\)", ")").replace("\\\\", "\\") for c in chunks)[:MAX_TEXT]
            a.kind, a.text = "text", text
            a.summary = f"PDF, {text.count(chr(12)) + 1} page(s) read, {len(text)} characters"
            if not text.strip():
                a.summary += " (no text layer: scanned PDF; OCR " + ("available" if self.ocr else "not installed") + ")"
            self._broker_hint(a, text)
        elif a.ext == "xml":
            self._xml(a, data)
        else:
            self._image(a, path, data)

    def _table_from_rows(self, a: Attachment, rows: list[list[str]]) -> None:
        rows = [r for r in rows if any(str(c).strip() for c in r)]
        if not rows:
            raise AttachmentError("parse_failed", "no rows found")
        head, body = [str(c)[:60] for c in rows[0]], rows[1:MAX_ROWS + 1]
        a.table = {"columns": head, "rows": [[str(c)[:200] for c in r] for r in body], "truncated": len(rows) - 1 > MAX_ROWS}
        a.kind = "table"
        a.summary = f"table, {len(body)} rows × {len(head)} columns" + (" (truncated)" if a.table["truncated"] else "")
        a.text = "\n".join([",".join(head)] + [",".join(r) for r in a.table["rows"][:200]])[:MAX_TEXT]
        self._broker_hint(a, " ".join(head))

    @staticmethod
    def _broker_hint(a: Attachment, text: str) -> None:
        low = text.lower()
        hits = [k for k in ("symbol", "quantity", "qty", "price", "commission", "settlement", "trade date", "book value", "market value") if k in low]
        if len(hits) >= 3:
            a.kind = "broker_statement" if a.kind in ("table", "text") else a.kind
            a.meta["broker_columns"] = hits
            a.summary += " · looks like a broker statement"

    def _xml(self, a: Attachment, data: bytes) -> None:
        from defusedxml import ElementTree as DET
        from defusedxml.common import DefusedXmlException
        try:
            root = DET.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
        except DefusedXmlException:
            raise AttachmentError("unsafe_xml", "XML with DTDs / entities is refused (XXE protection)")
        except DET.ParseError:
            raise AttachmentError("parse_failed", "XML is not well-formed")
        ns = root.tag.split("}")[0].strip("{") if root.tag.startswith("{") else ""
        a.meta["namespace"] = ns
        local = lambda e: e.tag.split("}")[-1]  # noqa: E731
        if "iso:std:iso:20022" in ns:
            msg = ns.rsplit(":", 1)[-1]
            a.kind = "iso20022"
            entries, bals = [], []
            for e in root.iter():
                t = local(e)
                if t == "Bal":
                    amt = next((x for x in e.iter() if local(x) == "Amt"), None)
                    cd = next((x.text for x in e.iter() if local(x) == "Cd"), "")
                    if amt is not None:
                        bals.append([cd, amt.text, amt.get("Ccy", "")])
                if t == "Ntry":
                    amt = next((x for x in e.iter() if local(x) == "Amt"), None)
                    ind = next((x.text for x in e.iter() if local(x) == "CdtDbtInd"), "")
                    bd = next((x for x in e.iter() if local(x) in ("BookgDt", "ValDt")), None)
                    dt = next((x.text for x in bd.iter() if local(x) in ("Dt", "DtTm")), "") if bd is not None else ""
                    info = next((x.text for x in e.iter() if local(x) in ("Ustrd", "AddtlNtryInf")), "")
                    entries.append([dt or "", ind, amt.text if amt is not None else "", amt.get("Ccy", "") if amt is not None else "", (info or "")[:120]])
            a.table = {"columns": ["date", "credit/debit", "amount", "ccy", "info"], "rows": entries[:MAX_ROWS], "truncated": len(entries) > MAX_ROWS}
            a.meta |= {"message": msg, "balances": bals[:10], "entries": len(entries)}
            a.summary = f"ISO 20022 {msg}: {len(entries)} entries, {len(bals)} balances"
            a.text = json.dumps({"message": msg, "balances": bals[:10], "entries": entries[:100]})[:MAX_TEXT]
        else:
            a.kind = "text"
            a.text = " ".join(t.strip() for t in root.itertext() if t.strip())[:MAX_TEXT]
            a.summary = f"XML <{local(root)}>, {len(list(root.iter()))} elements"

    def _image(self, a: Attachment, path: Path, data: bytes) -> None:
        try:
            from PIL import Image
            Image.MAX_IMAGE_PIXELS = 40_000_000           # decompression-bomb guard
            with Image.open(io.BytesIO(data)) as im:
                im.verify()
            with Image.open(io.BytesIO(data)) as im:
                a.meta = {"format": im.format, "width": im.width, "height": im.height, "mode": im.mode}
        except ImportError:
            import struct
            if data.startswith(b"\x89PNG\r\n\x1a\n") and len(data) >= 24 and data[12:16] == b"IHDR":
                w, h = struct.unpack(">II", data[16:24])
                fmt, mode = "PNG", "RGB"
            elif data.startswith((b"GIF87a", b"GIF89a")) and len(data) >= 10:
                w, h = struct.unpack("<HH", data[6:10])
                fmt, mode = "GIF", "P"
            elif data.startswith(b"\xff\xd8\xff"):
                w, h, fmt, mode = 1, 1, "JPEG", "RGB"
            elif data.startswith(b"RIFF") and data[8:12] == b"WEBP":
                w, h, fmt, mode = 1, 1, "WEBP", "RGB"
            else:
                raise AttachmentError("parse_failed", "image could not be verified")
            if w <= 0 or h <= 0 or w * h > 40_000_000:
                raise AttachmentError("parse_failed", "image dimensions invalid or exceed limit")
            a.meta = {"format": fmt, "width": int(w), "height": int(h), "mode": mode}
        a.kind = "image"
        a.summary = f"image {a.meta['format']} {a.meta['width']}×{a.meta['height']}"
        if self.ocr:
            try:
                out = subprocess.run(["tesseract", str(path), "-", "--psm", "3"], capture_output=True, timeout=30, check=False,
                                     env={"PATH": os.environ.get("PATH", "/usr/bin"), "OMP_THREAD_LIMIT": "1"})
                a.text = out.stdout.decode("utf-8", errors="replace")[:MAX_TEXT]
                a.summary += f" · OCR {len(a.text)} chars"
            except subprocess.TimeoutExpired:
                a.summary += " · OCR timed out"
        else:
            a.summary += " · no OCR installed (metadata only)"
